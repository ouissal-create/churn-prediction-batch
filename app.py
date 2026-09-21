import streamlit as st
import pandas as pd
import numpy as np
import pickle
import plotly.express as px
from pathlib import Path
import io

# ============================================================
# CONFIGURATION DE LA PAGE
# ============================================================
st.set_page_config(
    page_title="Churn Bancaire - Batch Prediction",
    layout="wide",
    initial_sidebar_state="expanded"
)

REQUIRED_COLUMNS = [
    'CreditScore', 'Age', 'Tenure', 'Balance', 'NumOfProducts',
    'HasCrCard', 'IsActiveMember', 'EstimatedSalary', 'Gender', 'Geography'
]

VALUE_RANGES = {
    'CreditScore': (300, 850),
    'Age': (18, 100),
}

SEUIL_RISQUE_ELEVE = 60
SEUIL_RISQUE_MODERE = 30

# Palette Mixte Bleu Royal / Orange / Bleu Ciel
COLORS = {
    'Élevé': '#ea580c',   # Orange foncé (Alerte)
    'Modéré': '#f59e0b',  # Orange ambré (Attention)
    'Faible': '#38bdf8'   # Bleu Ciel (Sécurité)
}

MODEL_PATH = Path("rf_model.pkl")
SCALER_PATH = Path("scaler.pkl")
FEATURES_PATH = Path("features_order.pkl")
GEO_CATEGORIES_PATH = Path("geography_categories.pkl")

# ============================================================
# CSS SUR MESURE (DESIGN SYSTEM MODERN & PRO)
# ============================================================
CUSTOM_CSS = """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');

    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
    }
    
    /* Arrière-plan global */
    .stApp {
        background-color: #f8fafc !important;
    }
    
    /* Sidebar styling */
    section[data-testid="stSidebar"] {
        background-color: #0f172a !important;
        border-right: 1px solid #1e293b;
    }
    
    section[data-testid="stSidebar"] * {
        color: #f8fafc !important;
    }
    
    /* Zone d'upload dans la Sidebar */
    div[data-testid="stFileUploader"] {
        background-color: #1e293b !important;
        border-radius: 12px;
        padding: 16px;
        border: 2px dashed #475569;
        transition: all 0.3s ease;
    }
    div[data-testid="stFileUploader"]:hover {
        border-color: #38bdf8;
    }
    
    div[data-testid="stFileUploader"] button[kind="secondary"] {
        background-color: #38bdf8 !important;
        color: #0f172a !important;
        border: none !important;
        font-weight: 600 !important;
        border-radius: 6px !important;
    }

    /* En-tête principal */
    .hero-header {
        background: linear-gradient(135deg, #0f172a 0%, #1e3a8a 50%, #1e40af 100%);
        color: white;
        padding: 2.5rem 2rem;
        border-radius: 20px;
        margin-bottom: 2rem;
        box-shadow: 0 10px 25px -5px rgba(15, 23, 42, 0.25);
        border: 1px solid rgba(255, 255, 255, 0.1);
    }
    
    .hero-header h1 {
        font-size: 2.2rem !important;
        font-weight: 800 !important;
        margin-bottom: 0.5rem !important;
        color: #ffffff !important;
    }
    
    .hero-header p {
        font-size: 1.05rem;
        color: #93c5fd;
        margin: 0;
    }

    /* Target Onglets Streamlit */
    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        background-color: #e2e8f0;
        padding: 6px;
        border-radius: 12px;
    }

    .stTabs [data-baseweb="tab"] {
        height: 45px;
        border-radius: 8px;
        font-weight: 600;
        color: #475569;
        border: none !important;
    }

    .stTabs [aria-selected="true"] {
        background-color: #ffffff !important;
        color: #1e3a8a !important;
        box-shadow: 0 2px 4px rgba(0,0,0,0.05);
    }

    /* Card Metrics personnalisées */
    .kpi-card {
        background-color: #ffffff;
        border-radius: 14px;
        padding: 1.25rem;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05), 0 2px 4px -1px rgba(0, 0, 0, 0.03);
        border: 1px solid #e2e8f0;
        position: relative;
        overflow: hidden;
    }
    
    .kpi-title {
        font-size: 0.875rem;
        font-weight: 600;
        color: #64748b;
        text-transform: uppercase;
        letter-spacing: 0.05em;
    }
    
    .kpi-value {
        font-size: 1.85rem;
        font-weight: 800;
        color: #0f172a;
        margin: 0.3rem 0;
    }
    
    .kpi-sub {
        font-size: 0.85rem;
        font-weight: 500;
    }
    
    .kpi-accent {
        position: absolute;
        top: 0;
        left: 0;
        width: 6px;
        height: 100%;
    }

    /* Recommandation Cards */
    .reco-card {
        background-color: #ffffff;
        border-radius: 12px;
        padding: 1.25rem;
        margin-bottom: 1rem;
        border-left: 5px solid;
        box-shadow: 0 2px 4px rgba(0,0,0,0.02);
    }

    /* Export Card */
    .export-card {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 14px;
        padding: 1.5rem;
        height: 100%;
        box-shadow: 0 4px 6px -1px rgba(0,0,0,0.03);
        display: flex;
        flex-direction: column;
        justify-content: space-between;
    }

    /* Titres des sous-sections */
    .section-header {
        font-size: 1.25rem;
        font-weight: 700;
        color: #0f172a;
        margin: 1.5rem 0 1rem 0;
        display: flex;
        align-items: center;
        gap: 0.5rem;
    }
</style>
"""

# ============================================================
# CHARGEMENT DES ARTEFACTS
# ============================================================
@st.cache_resource(show_spinner=False)
def charger_modeles():
    artefacts = {}
    for nom, chemin in [
        ("modele", MODEL_PATH), ("scaler", SCALER_PATH), ("ordre_colonnes", FEATURES_PATH)
    ]:
        if not chemin.exists():
            raise FileNotFoundError(f"Fichier requis introuvable : {chemin}")
        with open(chemin, 'rb') as f:
            artefacts[nom] = pickle.load(f)

    if GEO_CATEGORIES_PATH.exists():
        with open(GEO_CATEGORIES_PATH, 'rb') as f:
            artefacts["geo_categories"] = pickle.load(f)
    else:
        artefacts["geo_categories"] = ["France", "Germany", "Spain"]

    return artefacts["modele"], artefacts["scaler"], artefacts["ordre_colonnes"], artefacts["geo_categories"]

# ============================================================
# VALIDATION & PRETRAITEMENT
# ============================================================
def valider_colonnes(df: pd.DataFrame) -> list:
    return [col for col in REQUIRED_COLUMNS if col not in df.columns]

def valider_plages(df: pd.DataFrame) -> dict:
    anomalies = {}
    for col, (mini, maxi) in VALUE_RANGES.items():
        if col in df.columns:
            hors_plage = ((df[col] < mini) | (df[col] > maxi)).sum()
            if hors_plage > 0:
                anomalies[col] = int(hors_plage)
    return anomalies

def preprocess(df: pd.DataFrame, ordre_colonnes: list, geo_categories: list) -> tuple:
    df = df.copy()
    df['Geography_brute'] = df['Geography']
    df['Gender'] = df['Gender'].map({'Female': 0, 'Male': 1})
    df['Geography'] = pd.Categorical(df['Geography'], categories=geo_categories)
    df = pd.get_dummies(df, columns=['Geography'], drop_first=True)
    df['balance_to_salary'] = df['Balance'] / df['EstimatedSalary'].replace(0, np.nan)
    df['tenure_by_age'] = df['Tenure'] / df['Age'].replace(0, np.nan)
    df = df.fillna(0)
    X = df.reindex(columns=ordre_colonnes, fill_value=0)
    return df, X

def scorer_clients(X: pd.DataFrame, modele, scaler):
    X_scaled = scaler.transform(X)
    predictions = modele.predict(X_scaled)
    probabilites = modele.predict_proba(X_scaled)[:, 1] * 100
    return predictions, probabilites

def classifier_risque(proba: pd.Series) -> pd.Series:
    bins = [-float('inf'), SEUIL_RISQUE_MODERE, SEUIL_RISQUE_ELEVE, float('inf')]
    labels = ['Faible', 'Modéré', 'Élevé']
    return pd.cut(proba, bins=bins, labels=labels, right=False).astype(str)

@st.cache_data(show_spinner="Traitement analytique et prédictions en cours...")
def traiter_base_clients(df_clients, ordre_colonnes, geo_categories, _modele, _scaler):
    df_full, X = preprocess(df_clients, ordre_colonnes, geo_categories)
    predictions, probabilites = scorer_clients(X, _modele, _scaler)
    
    df_full['Prediction'] = predictions
    df_full['Probability_Churn'] = probabilites.round(2)
    df_full['Risque'] = classifier_risque(df_full['Probability_Churn'])
    df_full['Geography'] = df_full['Geography_brute']
    df_full['Gender_Label'] = df_full['Gender'].map({0: 'Femme', 1: 'Homme'})
    
    return df_full

@st.cache_data(show_spinner=False)
def _lire_csv(file): 
    return pd.read_csv(file)

@st.cache_data(show_spinner=False)
def _lire_excel(file): 
    return pd.read_excel(file)

def lire_fichier(uploaded_file):
    try:
        if uploaded_file.name.endswith('.csv'):
            return _lire_csv(uploaded_file)
        return _lire_excel(uploaded_file)
    except Exception as e:
        st.error(f"Erreur de lecture du fichier : {e}")
        return None

# ============================================================
# COMPOSANTS D'INTERFACE REFORMATÉS
# ============================================================
def afficher_kpi_card(title, value, sub_text, color, accent_color):
    st.markdown(f"""
    <div class="kpi-card">
        <div class="kpi-accent" style="background-color: {accent_color};"></div>
        <div class="kpi-title">{title}</div>
        <div class="kpi-value">{value}</div>
        <div class="kpi-sub" style="color: {color};">{sub_text}</div>
    </div>
    """, unsafe_allow_html=True)

def appliquer_theme_plotly(fig):
    fig.update_layout(
        paper_bgcolor='rgba(0,0,0,0)',
        plot_bgcolor='rgba(0,0,0,0)',
        font=dict(family="Inter, sans-serif", color="#334155"),
        margin=dict(l=20, r=20, t=40, b=20),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
    )
    fig.update_xaxes(showgrid=True, gridwidth=1, gridcolor='#f1f5f9')
    fig.update_yaxes(showgrid=True, gridwidth=1, gridcolor='#f1f5f9')
    return fig

# ============================================================
# VUES ET ONGLETS
# ============================================================
def afficher_vue_ensemble(df: pd.DataFrame):
    total = len(df)
    par_risque = df['Risque'].value_counts()
    
    n_eleve = int(par_risque.get('Élevé', 0))
    n_modere = int(par_risque.get('Modéré', 0))
    n_faible = int(par_risque.get('Faible', 0))

    # Grille de KPIs
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        afficher_kpi_card("Total Portefeuille", f"{total:,}".replace(",", " "), "Clients analysés", "#64748b", "#0f172a")
    with c2:
        afficher_kpi_card("Risque Élevé", f"{n_eleve:,}".replace(",", " "), f"{(n_eleve/total*100):.1f}% du total", COLORS['Élevé'], COLORS['Élevé'])
    with c3:
        afficher_kpi_card("Risque Modéré", f"{n_modere:,}".replace(",", " "), f"{(n_modere/total*100):.1f}% du total", COLORS['Modéré'], COLORS['Modéré'])
    with c4:
        afficher_kpi_card("Risque Faible", f"{n_faible:,}".replace(",", " "), f"{(n_faible/total*100):.1f}% du total", COLORS['Faible'], COLORS['Faible'])

    st.markdown("<div style='height: 20px;'></div>", unsafe_allow_html=True)

    # Graphiques principaux
    col_left, col_right = st.columns(2)
    
    with col_left:
        st.markdown("<div class='section-header'>Répartition du Niveau de Risque</div>", unsafe_allow_html=True)
        fig_pie = px.pie(
            df, names='Risque', 
            color='Risque', 
            color_discrete_map=COLORS,
            hole=0.5
        )
        fig_pie.update_traces(textinfo='percent+label', pull=[0.05, 0, 0])
        st.plotly_chart(appliquer_theme_plotly(fig_pie), use_container_width=True)

    with col_right:
        st.markdown("<div class='section-header'>Distribution des Probabilités de Churn</div>", unsafe_allow_html=True)
        fig_hist = px.histogram(
            df, x='Probability_Churn', nbins=30,
            color='Risque', color_discrete_map=COLORS,
            opacity=0.85
        )
        fig_hist.update_layout(xaxis_title='Probabilité de Churn (%)', yaxis_title='Nombre de clients')
        st.plotly_chart(appliquer_theme_plotly(fig_hist), use_container_width=True)

def afficher_segments(df: pd.DataFrame):
    st.markdown("<div class='section-header'>Analyse des Segments Clients</div>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    
    with c1:
        par_pays = df.groupby('Geography')['Probability_Churn'].mean().reset_index().sort_values('Probability_Churn', ascending=False)
        fig_geo = px.bar(
            par_pays, x='Geography', y='Probability_Churn',
            title='Probabilité moyenne de churn par Pays',
            color='Probability_Churn', color_continuous_scale='Oranges'
        )
        fig_geo.update_layout(xaxis_title="Pays", yaxis_title="Probabilité Moyenne (%)", coloraxis_showscale=False)
        st.plotly_chart(appliquer_theme_plotly(fig_geo), use_container_width=True)
        
    with c2:
        par_genre = df.groupby('Gender_Label')['Probability_Churn'].mean().reset_index()
        fig_gender = px.bar(
            par_genre, x='Gender_Label', y='Probability_Churn',
            title='Probabilité moyenne de churn par Genre',
            color='Gender_Label',
            color_discrete_sequence=['#38bdf8', '#818cf8']
        )
        fig_gender.update_layout(xaxis_title="Genre", yaxis_title="Probabilité Moyenne (%)", showlegend=False)
        st.plotly_chart(appliquer_theme_plotly(fig_gender), use_container_width=True)

def afficher_clients_risque(df: pd.DataFrame):
    st.markdown("<div class='section-header'>Exploration & Segmentation Détaillée</div>", unsafe_allow_html=True)
    
    # Filtres interactifs pour la table
    col_f1, col_f2 = st.columns([1, 3])
    with col_f1:
        filtre_risque = st.multiselect(
            "Filtrer par niveau de risque :",
            options=['Élevé', 'Modéré', 'Faible'],
            default=['Élevé', 'Modéré']
        )
    
    df_filtré = df[df['Risque'].isin(filtre_risque)] if filtre_risque else df
    df_trie = df_filtré.sort_values('Probability_Churn', ascending=False)
    
    colonnes_affichage = ['Probability_Churn', 'Risque'] + REQUIRED_COLUMNS[:-1] + ['Geography']
    colonnes_existantes = [col for col in colonnes_affichage if col in df_trie.columns]
    
    # Affichage du tableau de données stylisé Streamlit
    st.dataframe(
        df_trie[colonnes_existantes].head(100),
        use_container_width=True,
        height=400,
        column_config={
            "Probability_Churn": st.column_config.ProgressColumn(
                "Probabilité Churn (%)",
                format="%.2f%%",
                min_value=0,
                max_value=100,
            ),
            "Risque": st.column_config.TextColumn("Segment Risque")
        }
    )
    
    if len(df_trie) > 100:
        st.caption(f"Affichage restreint aux 100 premiers enregistrements (sur {len(df_trie)} clients correspondants). Consultez l'onglet **Export** pour obtenir l'intégralité.")

    st.markdown("<div style='height: 15px;'></div>", unsafe_allow_html=True)
    
    # Stats & Plan d'action côte à côte
    col_stat, col_reco = st.columns([1, 1.2])
    
    with col_stat:
        st.markdown("<div class='section-header'>Statistiques par Segment</div>", unsafe_allow_html=True)
        stats = df.groupby('Risque').agg({
            'Probability_Churn': ['count', 'mean', 'min', 'max']
        }).round(2)
        stats.columns = ['Nombre', 'Moyenne (%)', 'Min (%)', 'Max (%)']
        st.dataframe(stats, use_container_width=True)

    with col_reco:
        st.markdown("<div class='section-header'>Recommandations Stratégiques</div>", unsafe_allow_html=True)
        
        n_eleve = len(df[df['Risque'] == 'Élevé'])
        n_modere = len(df[df['Risque'] == 'Modéré'])
        n_faible = len(df[df['Risque'] == 'Faible'])

        if n_eleve > 0:
            st.markdown(f"""
            <div class="reco-card" style="border-color: {COLORS['Élevé']};">
                <strong style="color: {COLORS['Élevé']};">Risque Élevé ({n_eleve} clients)</strong>
                <ul style="margin: 0.5rem 0 0 0; padding-left: 1.2rem; font-size: 0.9rem; color: #334155;">
                    <li>Contact téléphonique prioritaire sous 48h par un conseiller expert.</li>
                    <li>Offres de rétention exclusives (réduction de frais, conditions préférentielles).</li>
                </ul>
            </div>
            """, unsafe_allow_html=True)
        
        if n_modere > 0:
            st.markdown(f"""
            <div class="reco-card" style="border-color: {COLORS['Modéré']};">
                <strong style="color: {COLORS['Modéré']};">Risque Modéré ({n_modere} clients)</strong>
                <ul style="margin: 0.5rem 0 0 0; padding-left: 1.2rem; font-size: 0.9rem; color: #334155;">
                    <li>Inclusion dans la campagne d'enquêtes de satisfaction automatisées.</li>
                    <li>Proposition de produits additionnels pour renforcer l'engagement.</li>
                </ul>
            </div>
            """, unsafe_allow_html=True)

        if n_faible > 0:
            st.markdown(f"""
            <div class="reco-card" style="border-color: {COLORS['Faible']};">
                <strong style="color: {COLORS['Faible']};">Risque Faible ({n_faible} clients)</strong>
                <ul style="margin: 0.5rem 0 0 0; padding-left: 1.2rem; font-size: 0.9rem; color: #334155;">
                    <li>Cibles privilégiées pour des actions de cross-selling (Cartes Gold, Investissements).</li>
                </ul>
            </div>
            """, unsafe_allow_html=True)

def gen_excel_bytes(df_input, style_fn, col_risk=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    style_fn(ws, df_input, col_risk)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def afficher_export(df: pd.DataFrame):
    """Export vers Excel avec mise en forme simple et sans couleurs."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    from openpyxl.utils import get_column_letter
    
    st.subheader("Export des Rapports")
    st.markdown("Téléchargez les rapports Excel au format standard.")
    
    # Styles simples sans couleurs
    FONT_HEADER = Font(name='Calibri', size=11, bold=True, color="000000") # Noir gras
    FONT_BODY = Font(name='Calibri', size=10, color="000000")             # Noir normal
    ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
    THIN_BORDER = Border(
        left=Side(style='thin', color="BFBFBF"),   # Bordure grise claire
        right=Side(style='thin', color="BFBFBF"),
        top=Side(style='thin', color="BFBFBF"),
        bottom=Side(style='thin', color="BFBFBF")
    )

    def style_sheet(ws, df_data, col_risque_name=None):
        for r_idx, row in enumerate(dataframe_to_rows(df_data, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if r_idx == 1:
                    # En-tête : fond blanc, texte noir gras
                    cell.font = FONT_HEADER
                    cell.alignment = ALIGN_CENTER
                    cell.border = THIN_BORDER
                else:
                    # Données : fond blanc, texte noir normal
                    cell.font = FONT_BODY
                    cell.border = THIN_BORDER

        # Ajuster largeur colonnes & Figer les volets & Ajouter filtres
        for col_cells in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                try:
                    if cell.value: max_length = max(max_length, len(str(cell.value)))
                except: pass
            ws.column_dimensions[col_letter].width = min(max_length + 4, 30)
        
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"

    # PRÉPARATION DES DONNÉES (une seule fois)
    df_exp = df.copy()
    df_exp['Risque'] = df_exp['Risque'].replace({'Élevé': 'Eleve', 'Modéré': 'Modere'})
    rename_map = {
        'Probability_Churn': 'Probabilite_Churn_Pct',
        'CreditScore': 'Score_Credit', 'Age': 'Age_Client',
        'Tenure': 'Anciennete_Annees', 'Balance': 'Solde_MAD',
        'NumOfProducts': 'Nb_Produits', 'HasCrCard': 'Carte_Bancaire',
        'IsActiveMember': 'App_Mobile_Active', 'EstimatedSalary': 'Revenu_Annuel_MAD',
        'Gender': 'Genre_Code', 'Geography': 'Pays_Residence',
        'Gender_Label': 'Genre'
    }
    df_exp = df_exp.rename(columns=rename_map)

    # GÉNÉRATION DES FICHIERS AVEC BARRE DE PROGRESSION
    progress_bar = st.progress(0, text="Génération des fichiers Excel en cours...")
    
    # Fichier 1 : Analyse Complète
    wb1 = Workbook(); ws1 = wb1.active; ws1.title = "Analyse_Complete"
    style_sheet(ws1, df_exp, col_risque_name='Risque')
    wb1.save('analyse_churn_complete.xlsx')
    progress_bar.progress(25, text="Fichier 1/4 généré : Analyse Complète")

    # Fichier 2 : Clients à Risque
    df_risk = df_exp[df_exp['Risque'].isin(['Eleve', 'Modere'])]
    wb2 = Workbook(); ws2 = wb2.active; ws2.title = "Clients_a_Risque"
    style_sheet(ws2, df_risk, col_risque_name='Risque')
    wb2.save('clients_a_risque.xlsx')
    progress_bar.progress(50, text="Fichier 2/4 généré : Clients à Risque")

    # Fichier 3 : Top 50 Prioritaire
    df_top = df_exp.sort_values('Probabilite_Churn_Pct', ascending=False).head(50).reset_index(drop=True)
    df_top.index += 1
    df_top.index.name = 'Rang_Priorite'
    df_top = df_top.reset_index()
    wb3 = Workbook(); ws3 = wb3.active; ws3.title = "Top_50_Prioritaire"
    style_sheet(ws3, df_top, col_risque_name='Risque')
    wb3.save('top_50_clients_risque.xlsx')
    progress_bar.progress(75, text="Fichier 3/4 généré : Top 50 Prioritaire")

    # Fichier 4 : Résumé Statistique
    stats = df.groupby('Risque').agg({
        'Probability_Churn': ['count', 'mean'], 'Balance': 'mean', 
        'Age': 'mean', 'NumOfProducts': 'mean'
    }).round(2)
    stats.columns = ['Nb_Clients', 'Proba_Moyenne', 'Solde_Moyen', 'Age_Moyen', 'Produits_Moyens']
    stats = stats.reset_index().rename(columns={'Risque': 'Niveau_Risque'})
    wb4 = Workbook(); ws4 = wb4.active; ws4.title = "Resume_Statistiques"
    style_sheet(ws4, stats)
    wb4.save('resume_statistiques.xlsx')
    progress_bar.progress(100, text="Tous les fichiers sont prêts !")

    # AFFICHAGE DES BOUTONS DE TÉLÉCHARGEMENT (instantané après génération)
    st.markdown("### Fichiers disponibles")
    
    c1, c2 = st.columns(2)
    
    with c1:
        st.markdown("**1. Analyse Complète**")
        st.caption("Tous les clients avec niveau de risque")
        with open('analyse_churn_complete.xlsx', 'rb') as f:
            st.download_button(
                label="Télécharger Analyse Complète",
                data=f.read(),
                file_name='analyse_churn_complete.xlsx',
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        
        st.markdown("---")
        
        st.markdown("**3. Top 50 Prioritaire**")
        st.caption("Clients les plus critiques numérotés par rang")
        with open('top_50_clients_risque.xlsx', 'rb') as f:
            st.download_button(
                label="Télécharger Top 50",
                data=f.read(),
                file_name='top_50_clients_risque.xlsx',
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    
    with c2:
        st.markdown("**2. Clients à Risque**")
        st.caption("Liste filtrée Eleve + Modere pour action immédiate")
        with open('clients_a_risque.xlsx', 'rb') as f:
            st.download_button(
                label="Télécharger Clients à Risque",
                data=f.read(),
                file_name='clients_a_risque.xlsx',
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        
        st.markdown("---")
        
        st.markdown("**4. Résumé Statistique**")
        st.caption("Indicateurs agrégés par segment de risque")
        with open('resume_statistiques.xlsx', 'rb') as f:
            st.download_button(
                label="Télécharger Résumé Stats",
                data=f.read(),
                file_name='resume_statistiques.xlsx',
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

def afficher_instructions():
    st.markdown("""
    <div style="background-color: #ffffff; border-radius: 16px; padding: 2.5rem; border: 1px solid #e2e8f0; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05);">
        <h3 style="color: #0f172a; margin-top: 0;">Bienvenue sur la plateforme de Prédiction de Churn</h3>
        <p style="color: #64748b;">Veuillez importer un fichier dans le panneau latéral pour lancer l'évaluation prédictive en batch.</p>
        <hr style="border: none; border-top: 1px solid #f1f5f9; margin: 1.5rem 0;">
        <h4 style="color: #1e293b;">Structure des données attendue :</h4>
        <p style="color: #475569; font-size: 0.95rem;">Le fichier transmis (CSV ou Excel) doit obligatoirement inclure les colonnes ci-dessous :</p>
        <div style="background-color: #f8fafc; padding: 1rem; border-radius: 8px; font-family: monospace; font-size: 0.88rem; color: #0f172a; border: 1px solid #e2e8f0;">
            CreditScore, Age, Tenure, Balance, NumOfProducts, HasCrCard, IsActiveMember, EstimatedSalary, Gender, Geography
        </div>
    </div>
    """, unsafe_allow_html=True)

# ============================================================
# MAIN
# ============================================================
def main():
    st.markdown(CUSTOM_CSS, unsafe_allow_html=True)

    # Banner Header
    st.markdown("""
    <div class="hero-header">
        <h1>Predictive Churn Analytics</h1>
        <p>Plateforme décisionnelle d'évaluation du risque d'attrition bancaire</p>
    </div>
    """, unsafe_allow_html=True)

    try:
        modele, scaler, ordre_colonnes, geo_categories = charger_modeles()
    except Exception as e:
        st.error(f"Erreur critique lors du chargement des artefacts ML : {e}")
        st.stop()

    # Sidebar Configuration
    st.sidebar.markdown("<h2 style='font-size: 1.2rem;'>Configuration</h2>", unsafe_allow_html=True)
    uploaded_file = st.sidebar.file_uploader(
        "Importer la base clients",
        type=['csv', 'xlsx', 'xls'],
        help=f"Colonnes requises : {', '.join(REQUIRED_COLUMNS)}"
    )

    if uploaded_file is None:
        afficher_instructions()
        return

    # Chargement
    df_clients = lire_fichier(uploaded_file)
    if df_clients is None: 
        return

    st.sidebar.success(f"Fichier chargé : {len(df_clients)} clients")
    
    with st.expander("Aperçu de la base importée", expanded=False):
        st.dataframe(df_clients.head(5), use_container_width=True)

    # Validations
    manquantes = valider_colonnes(df_clients)
    if manquantes:
        st.error(f"Colonnes manquantes dans votre fichier : {', '.join(manquantes)}")
        st.info(f"Structure requise : {', '.join(REQUIRED_COLUMNS)}")
        return

    anomalies = valider_plages(df_clients)
    if anomalies:
        detail = ", ".join(f"{col} : {n} ligne(s)" for col, n in anomalies.items())
        st.warning(f"Valeurs aberrantes détectées ({detail}).")

    # Scoring
    try:
        df_full = traiter_base_clients(df_clients, ordre_colonnes, geo_categories, modele, scaler)
    except Exception as e:
        st.error(f"Erreur d'exécution du modèle : {e}")
        return

    # Nav Tabs
    tab_vue, tab_segments, tab_risque, tab_export = st.tabs([
        "Vue d'ensemble", 
        "Segments", 
        "Clients à risque", 
        "Exporter les rapports"
    ])

    with tab_vue:
        afficher_vue_ensemble(df_full)
    with tab_segments:
        afficher_segments(df_full)
    with tab_risque:
        afficher_clients_risque(df_full)
    with tab_export:
        afficher_export(df_full)

if __name__ == "__main__":
    main()
